"""
Tests for weighted GDP calculation task - variant: range (max − min)
Verifies that the Excel formulas are correctly implemented.
"""

import re
import csv
import glob
from pathlib import Path

import pytest
from openpyxl import load_workbook

EXCEL_FILE = Path("/root/gdp.xlsx")
CSV_PATTERN = "/root/sheet.csv.*"
TOLERANCE = 0.5


def find_task_csv():
    csv_files = sorted(glob.glob(CSV_PATTERN))
    if not csv_files:
        return None
    wb = load_workbook(EXCEL_FILE, data_only=False)
    for idx, name in enumerate(wb.sheetnames):
        if "Task" in name:
            wb.close()
            expected_file = f"/root/sheet.csv.{idx}"
            if Path(expected_file).exists():
                return expected_file
            break
    wb.close()
    return csv_files[0] if csv_files else None


class TestGDPVariant:
    def test_excel_exists(self):
        assert EXCEL_FILE.exists(), f"Excel file not found: {EXCEL_FILE}"

    def test_task_sheet_has_data(self):
        wb = load_workbook(EXCEL_FILE, data_only=True)
        task_sheets = [s for s in wb.sheetnames if "Task" in s]
        assert task_sheets, "No Task sheet found"
        ws = wb[task_sheets[0]]
        wb.close()
        has_data = any(ws.cell(row=r, column=c).value is not None
                       for r in range(1, 50) for c in range(1, 20))
        assert has_data, "Task sheet appears empty"

    def test_net_exports_section_filled(self):
        wb = load_workbook(EXCEL_FILE, data_only=True)
        task_sheets = [s for s in wb.sheetnames if "Task" in s]
        assert task_sheets
        ws = wb[task_sheets[0]]
        wb.close()
        # Rows 35-40, columns H-L (8-12): net exports % GDP
        non_empty = sum(
            1 for r in range(35, 41) for c in range(8, 13)
            if ws.cell(row=r, column=c).value is not None
        )
        assert non_empty > 0, "Net exports % GDP cells (H35:L40) appear empty"

    def test_stat_formula_present(self):
        """Check that the correct formula type is used in the Task sheet."""
        wb = load_workbook(EXCEL_FILE, data_only=False)
        task_sheets = [s for s in wb.sheetnames if "Task" in s]
        assert task_sheets
        ws = wb[task_sheets[0]]
        wb.close()
        formula_found = False
        for row in ws.iter_rows():
            for cell in row:
                val = str(cell.value or "").upper()
                if re.search(r"MAX.*MIN|MIN.*MAX", val):
                    formula_found = True
                    break
            if formula_found:
                break
        assert formula_found, f"Expected a RANGE formula in Task sheet"

    def test_stat_result_is_numeric(self):
        wb = load_workbook(EXCEL_FILE, data_only=True)
        task_sheets = [s for s in wb.sheetnames if "Task" in s]
        assert task_sheets
        ws = wb[task_sheets[0]]
        wb.close()
        # The result should be somewhere in the task sheet as a numeric value
        numeric_vals = [
            cell.value for row in ws.iter_rows() for cell in row
            if isinstance(cell.value, (int, float))
        ]
        assert len(numeric_vals) > 0, "Task sheet should contain numeric result values"
